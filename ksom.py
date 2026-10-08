import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# KSOM
# 4 Inputs × 4 Features
# 2 Clusters
# Weight Matrix = 2 × 4
# ============================================================


# ------------------------------------------------------------
# 1. INPUT DATA
# ------------------------------------------------------------

X = np.array([
    [1, 0, 1, 0],   # 1010
    [1, 0, 0, 0],   # 1000
    [1, 1, 1, 1],   # 1111
    [0, 1, 1, 0]    # 0110
], dtype=float)


# ------------------------------------------------------------
# 2. INITIAL WEIGHTS
# ------------------------------------------------------------

W = np.array([
    [0.3, 0.5, 0.7, 0.2],   # Cluster 1
    [0.6, 0.5, 0.4, 0.2]    # Cluster 2
], dtype=float)


# ------------------------------------------------------------
# 3. PARAMETERS
# ------------------------------------------------------------

learning_rate = 0.6
epsilon = 0.001
max_iterations = 100


# ------------------------------------------------------------
# 4. STORE RESULTS
# ------------------------------------------------------------

results = []


# ============================================================
# 5. KSOM TRAINING
# ============================================================

for iteration in range(1, max_iterations + 1):

    old_W = W.copy()

    print("\n================================")
    print("ITERATION:", iteration)
    print("================================")

    for input_no, x in enumerate(X, start=1):

        # Euclidean distance
        distance_1 = np.sqrt(
            np.sum((x - W[0]) ** 2)
        )

        distance_2 = np.sqrt(
            np.sum((x - W[1]) ** 2)
        )

        # Find winner
        if distance_1 < distance_2:
            winner = 0
        else:
            winner = 1

        # Old winner weight
        old_weight = W[winner].copy()

        # Weight update
        W[winner] = (
            W[winner]
            + learning_rate * (x - W[winner])
        )

        # Weight change
        change = np.max(
            np.abs(W[winner] - old_weight)
        )

        # Print
        print("\nInput:", input_no)
        print("X:", x.astype(int))
        print("Distance C1:", distance_1)
        print("Distance C2:", distance_2)
        print("Winner: Cluster", winner + 1)
        print("Old Weight:", old_weight)
        print("New Weight:", W[winner])

        # Store
        results.append({
            "Iteration": iteration,
            "Input": input_no,
            "Input_Vector": "".join(str(int(v)) for v in x),

            "Distance_C1": distance_1,
            "Distance_C2": distance_2,

            "Winner": winner + 1,

            "Learning_Rate": learning_rate,

            "Old_W1": old_weight[0],
            "Old_W2": old_weight[1],
            "Old_W3": old_weight[2],
            "Old_W4": old_weight[3],

            "New_W1": W[winner][0],
            "New_W2": W[winner][1],
            "New_W3": W[winner][2],
            "New_W4": W[winner][3],

            "Weight_Change": change
        })

    # --------------------------------------------------------
    # Stopping condition
    # --------------------------------------------------------

    weight_change = np.max(
        np.abs(W - old_W)
    )

    print("\nMaximum Weight Change:", weight_change)

    if weight_change < epsilon:

        print("\n>>> TRAINING STOPPED <<<")
        print("Reason: Weight change < epsilon")

        break


# ============================================================
# 6. FINAL WEIGHTS
# ============================================================

final_weights = pd.DataFrame(
    W,
    index=["Cluster 1", "Cluster 2"],
    columns=["W1", "W2", "W3", "W4"]
)


# ============================================================
# 7. FINAL CLUSTER ASSIGNMENT
# ============================================================

clusters = []

for input_no, x in enumerate(X, start=1):

    distance_1 = np.sqrt(
        np.sum((x - W[0]) ** 2)
    )

    distance_2 = np.sqrt(
        np.sum((x - W[1]) ** 2)
    )

    if distance_1 < distance_2:
        cluster = 1
    else:
        cluster = 2

    clusters.append({
        "Input": input_no,
        "Vector": "".join(str(int(v)) for v in x),
        "Distance_C1": distance_1,
        "Distance_C2": distance_2,
        "Cluster": cluster
    })

cluster_df = pd.DataFrame(clusters)


# ============================================================
# 8. CREATE EXCEL
# ============================================================

file_name = "KSOM_Result.xlsx"

with pd.ExcelWriter(
    file_name,
    engine="openpyxl"
) as writer:

    # Input Data
    pd.DataFrame(
        X.astype(int),
        columns=["X1", "X2", "X3", "X4"]
    ).to_excel(
        writer,
        sheet_name="Input Data",
        index=False
    )

    # Iterations
    pd.DataFrame(
        results
    ).to_excel(
        writer,
        sheet_name="Iterations",
        index=False
    )

    # Final Weights
    final_weights.to_excel(
        writer,
        sheet_name="Final Weights"
    )

    # Clusters
    cluster_df.to_excel(
        writer,
        sheet_name="Clusters",
        index=False
    )


# ============================================================
# 9. EXCEL FORMATTING
# ============================================================

# Load workbook
from openpyxl import load_workbook

wb = load_workbook(file_name)


# ------------------------------------------------------------
# Colors
# ------------------------------------------------------------

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

title_fill = PatternFill(
    fill_type="solid",
    fgColor="17365D"
)

subheader_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

bold_font = Font(
    bold=True
)

thin_side = Side(
    style="thin",
    color="B7B7B7"
)

border = Border(
    left=thin_side,
    right=thin_side,
    top=thin_side,
    bottom=thin_side
)


# ------------------------------------------------------------
# Format each worksheet
# ------------------------------------------------------------

for ws in wb.worksheets:

    # Header row
    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        cell.border = border

    # All cells
    for row in ws.iter_rows():

        for cell in row:

            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    # Column width
    for column in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            30
        )

    # Freeze first row
    ws.freeze_panes = "A2"

    # Header height
    ws.row_dimensions[1].height = 25


# ============================================================
# 10. SPECIAL FORMATTING
# ============================================================

# Final Weights sheet
ws = wb["Final Weights"]

for cell in ws[1]:
    cell.fill = title_fill
    cell.font = header_font

# Make cluster names bold
for cell in ws["A"]:
    cell.font = bold_font


# Iterations sheet
ws = wb["Iterations"]

# Format decimal numbers
for row in ws.iter_rows(min_row=2):

    for cell in row:

        if isinstance(cell.value, float):
            cell.number_format = "0.0000"


# Final Weights decimal format
ws = wb["Final Weights"]

for row in ws.iter_rows(min_row=2):

    for cell in row:

        if isinstance(cell.value, float):
            cell.number_format = "0.0000"


# Clusters decimal format
ws = wb["Clusters"]

for row in ws.iter_rows(min_row=2):

    for cell in row:

        if isinstance(cell.value, float):
            cell.number_format = "0.0000"


# ------------------------------------------------------------
# Save formatted workbook
# ------------------------------------------------------------

wb.save(file_name)


# ============================================================
# 11. FINAL OUTPUT
# ============================================================

print("\n================================")
print("KSOM COMPLETED")
print("================================")

print("\nFinal Weight Matrix:")
print(final_weights)

print("\nCluster Assignment:")
print(cluster_df)

print("\nIterations:", iteration)

print("\nExcel file created:")
print(file_name)